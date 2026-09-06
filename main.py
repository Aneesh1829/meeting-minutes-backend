import os
import io
import uuid
from fastapi import FastAPI, UploadFile, File, BackgroundTasks, Header, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from supabase import create_client
from groq import Groq
from dotenv import load_dotenv
import markdown as md_lib
from xhtml2pdf import pisa
from docx import Document

load_dotenv()

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

supabase = create_client(os.getenv("SUPABASE_URL"), os.getenv("SUPABASE_SERVICE_KEY"))
groq_client = Groq(api_key=os.getenv("GROQ_API_KEY"))
BUCKET = os.getenv("SUPABASE_BUCKET")
APP_API_KEY = os.getenv("APP_API_KEY")

MAX_FILE_SIZE_MB = 25  # matches Groq's free-tier Whisper limit
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

SYSTEM_MESSAGE = """
You produce minutes of meetings from transcripts, with summary, key discussion points,
takeaways and action items with owners, in markdown format without code blocks.
"""


def verify_api_key(x_api_key: str = Header(None)):
    if not APP_API_KEY:
        # No key configured on the server — skip the check (dev convenience only)
        return
    if x_api_key != APP_API_KEY:
        raise HTTPException(status_code=401, detail="Invalid or missing API key")


def process_meeting(meeting_id: str, storage_path: str, local_audio_path: str):
    try:
        with open(local_audio_path, "rb") as audio_file:
            transcription = groq_client.audio.transcriptions.create(
                file=(local_audio_path, audio_file.read()),
                model="whisper-large-v3-turbo",
            )
        transcript_text = transcription.text

        user_prompt = f"""
Below is a transcript of a meeting.
Please write minutes in markdown without code blocks, including:
- a summary with attendees, location and date (if mentioned)
- discussion points
- takeaways
- action items with owners

Transcription:
{transcript_text}
"""
        completion = groq_client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[
                {"role": "system", "content": SYSTEM_MESSAGE},
                {"role": "user", "content": user_prompt},
            ],
            max_tokens=2000,
        )
        minutes_text = completion.choices[0].message.content

        supabase.table("meetings").update({
            "status": "done",
            "transcript": transcript_text,
            "minutes": minutes_text,
        }).eq("id", meeting_id).execute()

    except Exception as e:
        supabase.table("meetings").update({"status": "failed"}).eq("id", meeting_id).execute()
        print(f"Error processing meeting {meeting_id}: {e}")

    finally:
        if os.path.exists(local_audio_path):
            os.remove(local_audio_path)


@app.get("/health")
async def health():
    # Lightweight — no Supabase/Groq calls — used by the frontend to detect
    # a Render free-tier cold start without triggering real work.
    return {"status": "ok"}


@app.post("/meetings")
async def upload_meeting(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    _: None = None,
    x_api_key: str = Header(None),
):
    verify_api_key(x_api_key)

    file_bytes = await file.read()

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        size_mb = round(len(file_bytes) / (1024 * 1024), 1)
        return JSONResponse(
            status_code=413,
            content={"error": f"File is {size_mb}MB — the limit is {MAX_FILE_SIZE_MB}MB."}
        )

    meeting_id = str(uuid.uuid4())
    storage_path = f"{meeting_id}_{file.filename}"

    supabase.storage.from_(BUCKET).upload(storage_path, file_bytes)

    local_audio_path = f"temp_{storage_path}"
    with open(local_audio_path, "wb") as f:
        f.write(file_bytes)

    supabase.table("meetings").insert({
        "id": meeting_id,
        "filename": storage_path,
        "status": "processing",
    }).execute()

    background_tasks.add_task(process_meeting, meeting_id, storage_path, local_audio_path)

    return {"meeting_id": meeting_id, "status": "processing"}


@app.get("/meetings/{meeting_id}")
async def get_meeting(meeting_id: str, x_api_key: str = Header(None)):
    verify_api_key(x_api_key)
    result = supabase.table("meetings").select("*").eq("id", meeting_id).execute()
    if not result.data:
        return {"error": "not found"}
    return result.data[0]


@app.get("/meetings/{meeting_id}/download/pdf")
async def download_pdf(meeting_id: str, x_api_key: str = Header(None)):
    verify_api_key(x_api_key)
    result = supabase.table("meetings").select("*").eq("id", meeting_id).execute()
    if not result.data or result.data[0]["status"] != "done":
        return {"error": "not ready"}

    minutes_md = result.data[0]["minutes"]
    html = md_lib.markdown(minutes_md, extensions=["tables"])

    pdf_buffer = io.BytesIO()
    pisa.CreatePDF(io.StringIO(html), dest=pdf_buffer)
    pdf_buffer.seek(0)

    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=minutes_{meeting_id}.pdf"}
    )


@app.get("/meetings/{meeting_id}/download/docx")
async def download_docx(meeting_id: str, x_api_key: str = Header(None)):
    verify_api_key(x_api_key)
    result = supabase.table("meetings").select("*").eq("id", meeting_id).execute()
    if not result.data or result.data[0]["status"] != "done":
        return {"error": "not ready"}

    minutes_md = result.data[0]["minutes"]
    doc = Document()

    for line in minutes_md.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=minutes_{meeting_id}.docx"}
    )